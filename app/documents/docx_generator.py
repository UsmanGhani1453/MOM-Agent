from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.schemas import MeetingMinutes


def _bullets(doc: Document, items: list[str]) -> None:
    if not items:
        p = doc.add_paragraph("None recorded.")
        p.runs[0].italic = True
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def generate_docx(minutes: MeetingMinutes, out_path: Path, fallback_title: str = "Meeting") -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading(minutes.meeting_title or fallback_title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sub = doc.add_paragraph()
    r = sub.add_run("Minutes of Meeting" + (f"  |  {minutes.date}" if minutes.date else ""))
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    doc.add_heading("Participants", level=1)
    _bullets(doc, minutes.participants)

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(minutes.summary or "No summary available.")

    doc.add_heading("Discussion Points", level=1)
    _bullets(doc, minutes.discussion_points)

    doc.add_heading("Decisions", level=1)
    _bullets(doc, minutes.decisions)

    doc.add_heading("Action Items", level=1)
    if minutes.action_items:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        for cell, text in zip(table.rows[0].cells, ("Task", "Responsible", "Status")):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for a in minutes.action_items:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = a.task, a.responsible_person, a.status
    else:
        p = doc.add_paragraph("None recorded.")
        p.runs[0].italic = True

    doc.add_heading("Important Topics", level=1)
    _bullets(doc, minutes.important_topics)

    doc.add_heading("Next Steps", level=1)
    _bullets(doc, minutes.next_steps)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
